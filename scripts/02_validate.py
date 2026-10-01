"""Compare the two transcription passes of a province, run the arithmetic checks, write the long dataset.

Usage:
  .venv/bin/python scripts/02_validate.py Torino           compare passes, run checks, write outputs
  .venv/bin/python scripts/02_validate.py Torino --audit   also draw the random audit sheet (never overwrites)
  .venv/bin/python scripts/02_validate.py Torino --score   score the audit sheet filled in by a person

Inputs  work/<prov>/pass1/*.csv, pass2/*.csv   transcriptions (values as printed)
        work/<prov>/resolutions.csv            optional: cell values decided after re-reading a zoomed crop
        work/<prov>/blocks.csv                 block -> PDF page
Outputs output/tav10_<prov>_long.csv, output/flags_<prov>.csv, output/validation_<prov>.md
"""
import csv
import math
import random
import re
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
DASHES = {"—", "–", "-", "−"}
TOL = 0.005                      # superficie is printed with 2 decimals
VARS = ("az", "sup")
N_AUDIT = 200


# ---------------------------------------------------------------- parsing
def norm_raw(s):
    s = (s or "").strip().replace(" ", "")
    return "—" if s in DASHES else s


def parse(raw, var):
    """Printed string -> number. Returns (value, problem)."""
    s = norm_raw(raw)
    if s == "":
        return math.nan, "empty"
    if s == "—":
        return 0, None
    if "?" in s:
        return math.nan, "unreadable"
    if var == "az":
        if not re.fullmatch(r"\d{1,3}(\.\d{3})*", s):
            return math.nan, f"bad format '{s}'"
        return int(s.replace(".", "")), None
    if not re.fullmatch(r"\d{1,3}(\.\d{3})*,\d{2}", s):
        return math.nan, f"bad format '{s}'"
    return float(s.replace(".", "").replace(",", ".")), None


def read_pass(folder):
    """{block: {"ra_num":..,"zona":.., "cells": {(class_code, var): raw}}}; empty templates are skipped."""
    out = {}
    for f in sorted(folder.glob("*.csv")):
        rows = list(csv.reader(open(f)))
        meta = {r[0]: r[2].strip() for r in rows[1:3]}
        cells = {}
        for r in rows[3:]:
            cells[(int(r[0]), "az")], cells[(int(r[0]), "sup")] = r[2], r[3]
        if all(v.strip() == "" for v in cells.values()):
            continue
        out[f.stem] = {"ra_num": meta.get("ra_num", ""), "zona": meta.get("zona", ""), "cells": cells}
    return out


# ---------------------------------------------------------------- main validation
def validate(prov):
    work, outdir = ROOT / "work" / prov.lower(), ROOT / "output"
    outdir.mkdir(exist_ok=True)
    book = next(b for b in csv.DictReader(open(ROOT / "config" / "books.csv")) if b["provincia"] == prov)
    classes = pd.read_csv(ROOT / "config" / "classes.csv").set_index("class_code")
    pages = {r["block"]: r["pdf_page"] for r in csv.DictReader(open(work / "blocks.csv"))}
    p1, p2 = read_pass(work / "pass1"), read_pass(work / "pass2")
    resolutions = {}
    if (work / "resolutions.csv").exists():
        for r in csv.DictReader(open(work / "resolutions.csv")):
            if r["var"] in ("ra_num", "zona"):             # block metadata, class_code left blank
                resolutions[(r["block"], r["var"])] = (r["value_raw"], r["note"])
            else:
                resolutions[(r["block"], int(r["class_code"]), r["var"])] = (r["value_raw"], r["note"])

    blocks = sorted(p1)
    flags, rows = [], []
    pass2_complete = set(p2) >= set(p1)

    def flag(check, block, cc, var, detail):
        flags.append(dict(check=check, block=block, class_code=cc, var=var, pdf_page=pages.get(block, ""),
                          class_label=classes.loc[cc, "label"] if cc else "", detail=detail))

    for b in blocks:
        for key in ("ra_num", "zona"):
            if (b, key) in resolutions:
                p1[b][key] = resolutions[(b, key)][0]
            elif b in p2 and p1[b][key] != p2[b][key]:
                flag("pass disagreement", b, None, key, f"pass1 '{p1[b][key]}' vs pass2 '{p2[b][key]}'")
        for cc in classes.index:
            rec = dict(block=b, class_code=cc)
            for var in VARS:
                r1 = norm_raw(p1[b]["cells"][(cc, var)])
                r2 = norm_raw(p2[b]["cells"][(cc, var)]) if b in p2 else None
                if (b, cc, var) in resolutions:
                    raw, status = norm_raw(resolutions[(b, cc, var)][0]), "resolved"
                elif r2 is None:
                    raw, status = r1, "single pass"
                elif r1 == r2:
                    raw, status = r1, "agree"
                else:
                    raw, status = r1, "disagree"
                    flag("pass disagreement", b, cc, var, f"pass1 '{r1}' vs pass2 '{r2}'")
                val, problem = parse(raw, var)
                if problem:
                    flag("unparseable", b, cc, var, problem)
                rec.update({var: val, f"{var}_raw": raw, f"{var}_pass1": r1, f"{var}_pass2": r2,
                            f"{var}_status": status})
            rows.append(rec)
    df = pd.DataFrame(rows)

    # (a) within each block: sum of the 34 classes = TOTALE row
    col_fail = set()
    for b, g in df.groupby("block"):
        g = g.set_index("class_code")
        for var in VARS:
            s, t = g.loc[1:34, var].sum(), g.loc[35, var]
            if pd.isna(s) or pd.isna(t) or abs(s - t) > TOL:
                col_fail.add((b, var))
                flag("sum classes ≠ TOTALE", b, 35, var, f"sum of classes {s:,.2f} vs TOTALE {t:,.2f}")

    # (b)+(c) across regioni agrarie: sum of RAs = Tav. 1, class by class (incl. TOTALE)
    row_fail = set()
    ra = df[df.block.str.startswith("RA")]
    if "TAV1" in p1:
        tav1 = df[df.block == "TAV1"].set_index("class_code")
        for cc in classes.index:
            for var in VARS:
                s, t = ra.loc[ra.class_code == cc, var].sum(), tav1.loc[cc, var]
                if pd.isna(s) or pd.isna(t) or abs(s - t) > TOL:
                    row_fail.add((cc, var))
                    flag("sum RA ≠ Tav. 1", "TAV1", cc, var, f"sum of RAs {s:,.2f} vs Tav. 1 {t:,.2f}")

    # (d)+(e) plausibility of each class cell
    for r in df.itertuples():
        if r.class_code == 35:
            continue
        if r.class_code == 1:
            if r.sup != 0:
                flag("senza terreno with superficie", r.block, 1, "sup", f"sup {r.sup}")
            continue
        if (r.az == 0) != (r.sup == 0):
            flag("zero mismatch", r.block, r.class_code, "az/sup", f"az {r.az} sup {r.sup}")
            continue
        if r.az and r.az > 0:
            lo, hi = classes.loc[r.class_code, "lower_ha"], classes.loc[r.class_code, "upper_ha"]
            mean = r.sup / r.az
            if mean < lo - TOL or (not pd.isna(hi) and mean > hi + TOL):
                flag("mean outside class", r.block, r.class_code, "az/sup",
                     f"{r.sup:,.2f} ha / {r.az} = {mean:.3f} ha, class {lo}–{hi}")

    # localisation: a misread class cell breaks both its block sum and its Tav. 1 sum
    for b, var in col_fail:
        for cc, var2 in row_fail:
            if var == var2 and cc != 35:
                flag("suspect cell", b, cc, var, "fails both block sum and Tav. 1 sum")

    # ---------------------------------------------------------------- outputs
    meta = {b: p1[b] for b in blocks}
    long = df[df.block.str.startswith("RA")].copy()
    long.insert(0, "provincia", prov)
    long.insert(1, "sigla", book["sigla"])
    long.insert(2, "cod_prov", book["cod_prov"])
    long.insert(3, "reg_agr", long.block.map(lambda b: int(meta[b]["ra_num"]) if meta[b]["ra_num"].isdigit() else int(b[2:])))
    long.insert(4, "id", book["sigla"] + long.reg_agr.astype(str))
    long.insert(5, "zona", long.block.map(lambda b: meta[b]["zona"]))
    long.insert(7, "class_label", long.class_code.map(classes["label"]))
    long.insert(8, "suffix", long.class_code.map(classes["suffix"]))
    long["pdf_page"] = long.block.map(pages)
    long.to_csv(outdir / f"tav10_{prov.lower()}_long.csv", index=False)
    df[df.block == "TAV1"].to_csv(outdir / f"tav1_{prov.lower()}_long.csv", index=False)

    fl = pd.DataFrame(flags, columns=["check", "block", "class_code", "var", "pdf_page", "class_label", "detail"])
    fl.to_csv(outdir / f"flags_{prov.lower()}.csv", index=False)

    n_cells = len(df) * 2
    status = pd.concat([df[f"{v}_status"] for v in VARS]).value_counts()
    lines = [f"# Validation — {prov}", "",
             f"- blocks: {len(blocks)} ({sum(b.startswith('RA') for b in blocks)} regioni agrarie + Tav. 1)",
             f"- cells: {n_cells}",
             f"- pass 2 complete: {pass2_complete}",
             "- cell status: " + ", ".join(f"{k} {v}" for k, v in status.items()), "",
             "| check | flags |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in fl.check.value_counts().items()] or ["| (none) | 0 |"]
    (outdir / f"validation_{prov.lower()}.md").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))
    if len(fl):
        print("\nflags:")
        print(fl.to_string(index=False, max_colwidth=60))
    return long


# ---------------------------------------------------------------- audit sheet
def draw_audit(prov, long):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    path = ROOT / "output" / f"audit_{prov.lower()}.xlsx"
    if path.exists():
        print(f"\n{path.name} already exists: not overwritten")
        return
    rng = random.Random(1961)
    cells = [(r.id, r.block, r.pdf_page, r.class_label, var, getattr(r, f"{var}_raw"))
             for r in long.itertuples() for var in VARS]
    by_block = {}
    for c in cells:
        by_block.setdefault(c[1], []).append(c)
    per_block = N_AUDIT // len(by_block) + 1
    sample = [c for b in sorted(by_block) for c in rng.sample(by_block[b], per_block)]
    sample = sorted(rng.sample(sample, N_AUDIT), key=lambda c: (c[1], cells.index(c)))

    wb = Workbook()
    ws = wb.active
    ws.title = "audit"
    head = ["n", "id", "blocco", "pagina PDF", "classe (ettari)", "colonna TOTALE", "valore trascritto",
            "corretto? (si/no)", "valore corretto (se no)", "note"]
    ws.append(head)
    for i, (id_, b, page, label, var, raw) in enumerate(sample, 1):
        ws.append([i, id_, b, int(page), label, "Aziende" if var == "az" else "Superficie", raw, "", "", ""])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for col, w in zip("ABCDEFGHIJ", (5, 7, 8, 11, 20, 15, 18, 17, 22, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"si,no"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{N_AUDIT + 1}")
    ws.freeze_panes = "A2"

    info = wb.create_sheet("istruzioni")
    for line in [
        f"Audit della digitalizzazione — Tav. 10, provincia di {prov} ({N_AUDIT} celle estratte a caso)",
        "",
        "1. Aprire il PDF della provincia alla 'pagina PDF' indicata (numero di pagina del file, non quello stampato).",
        "2. Trovare il blocco 'REGIONE AGRARIA n' (colonna 'id': TO1 = regione agraria 1) e la riga della 'classe'.",
        "3. Guardare solo le colonne TOTALE (Aziende o Superficie, come indicato).",
        "4. Confrontare con 'valore trascritto': scrivere si se identico, no se diverso.",
        "5. Se no, scrivere nella colonna successiva il valore come stampato (es. 1.451,79 oppure —).",
        "",
        "Il trattino — significa zero/nessuna azienda.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110
    info["A1"].font = Font(bold=True)
    for r in info.iter_rows():
        r[0].alignment = Alignment(wrap_text=False)
    wb.save(path)
    print(f"\naudit sheet written: {path}")


def upper_95(k, n):
    """One-sided 95% Clopper-Pearson upper bound for a binomial proportion."""
    if k >= n:
        return 1.0
    cdf = lambda p: sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))
    lo, hi = k / n, 1.0
    for _ in range(60):
        mid = (lo + hi) / 2
        lo, hi = (mid, hi) if cdf(mid) > 0.05 else (lo, mid)
    return hi


def score_audit(prov):
    from openpyxl import load_workbook
    path = ROOT / "output" / f"audit_{prov.lower()}.xlsx"
    ws = load_workbook(path).active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    done = [r for r in rows if str(r[7] or "").strip().lower() in ("si", "sì", "no")]
    errors = [r for r in done if str(r[7]).strip().lower() == "no"]
    n, k = len(done), len(errors)
    lines = [f"# Audit result — {prov}", "",
             f"- cells checked: {n} of {len(rows)}",
             f"- errors found: {k}",
             f"- error rate: {k / n:.2%}" if n else "- error rate: n/a",
             f"- 95% upper bound on the error rate: {upper_95(k, n):.2%}" if n else ""]
    if errors:
        lines += ["", "| id | classe | colonna | trascritto | corretto | note |", "|---|---|---|---|---|---|"]
        lines += [f"| {r[1]} | {r[4]} | {r[5]} | {r[6]} | {r[8]} | {r[9] or ''} |" for r in errors]
    (ROOT / "output" / f"audit_result_{prov.lower()}.md").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))


if __name__ == "__main__":
    args = sys.argv[1:]
    prov = next((a for a in args if not a.startswith("--")), "Torino")
    if "--score" in args:
        score_audit(prov)
    else:
        long = validate(prov)
        if "--audit" in args:
            draw_audit(prov, long)
